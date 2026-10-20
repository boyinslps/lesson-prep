# -*- coding: utf-8 -*-
"""
（維護用，老師平常不用跑）把主辦單位的《小學-作品說明書》.odt 轉成帶填寫標記的 Word 範本 作品說明書_範本.docx。
需要 LibreOffice（soffice）與 python-docx。主辦單位改版時：換掉 odt、跑一次、再檢查 產生說明書.py 的標記對照。
    python 建立說明書範本.py  小學-作品說明書115.9.19.odt
"""
import sys, os, re, subprocess, tempfile, zipfile, shutil
import docx

src = os.path.abspath(sys.argv[1])
here = os.path.dirname(os.path.abspath(__file__))
tmp = tempfile.mkdtemp()
subprocess.run(['soffice', '--headless', '--convert-to', 'docx', '--outdir', tmp, src], check=True, stdout=subprocess.DEVNULL)
d = docx.Document(os.path.join(tmp, os.path.splitext(os.path.basename(src))[0] + '.docx'))


def settext(p, text, plain=False):
    p.runs[0].text = text
    if plain:
        p.runs[0].bold = False
    for r in p.runs[1:]:
        r._r.getparent().remove(r._r)


T = d.tables
settext(T[0].cell(0, 0).paragraphs[0], '{{idea}}', plain=True)
for i, p in enumerate(T[1].cell(1, 0).paragraphs):          # 六章勾選：把項目符號換成可填的方框
    numpr = p._p.pPr.find('{http://schemas.openxmlformats.org/wordprocessingml/2006/main}numPr')
    if numpr is not None:
        p._p.pPr.remove(numpr)
    p.runs[0].text = '{{c%d}} ' % (i + 1) + p.runs[0].text
    p.paragraph_format.left_indent = docx.shared.Cm(1.2)
    p.paragraph_format.first_line_indent = docx.shared.Cm(-0.6)
settext(T[1].cell(3, 0).paragraphs[0], '{{chapDesc}}', plain=True)
p = T[2].cell(0, 0).paragraphs[0]
p.runs[3].text = '\t\t{{ey}} '
p.runs[5].text = '\t\t{{en}} '
settext(T[2].cell(2, 0).paragraphs[0], '{{found}}', plain=True)
settext(T[2].cell(4, 0).paragraphs[0], '{{eUse}}', plain=True)
settext(T[3].cell(0, 0).paragraphs[0], '{{disclose}}', plain=True)
c = T[4].cell(1, 0)
settext(c.paragraphs[0], '{{te}} 因材網 e度\t{{tc}} Canva\t{{tp}} PowerPoint\t{{tw}} Word')
settext(c.paragraphs[1], '{{to}} 其他：{{toolsOther}}')
settext(T[4].cell(3, 0).paragraphs[0], '{{learned}}', plain=True)
mid = os.path.join(tmp, 'mid.docx')
d.save(mid)

# 拿掉 LibreOffice 轉檔時內嵌的字型（約 3 MB，Word 用不到）
out = os.path.join(here, '作品說明書_範本.docx')
with zipfile.ZipFile(mid) as zi, zipfile.ZipFile(out, 'w', zipfile.ZIP_DEFLATED) as zo:
    for it in zi.infolist():
        if it.filename.startswith('word/fonts/') or it.filename == 'word/_rels/fontTable.xml.rels':
            continue
        data = zi.read(it.filename)
        if it.filename == 'word/fontTable.xml':
            data = re.sub(rb'<w:embed\w+ [^>]*/>', b'', data)
        if it.filename == 'word/settings.xml':
            data = data.replace(b'<w:embedTrueTypeFonts/>', b'')
        zo.writestr(it.filename, data)
shutil.rmtree(tmp)
print('已產生', out, os.path.getsize(out) // 1024, 'KB')
